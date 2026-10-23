"""Customer reports — catalogue, generation, export and scheduled email delivery.

Every report is a function of (group, period, optional vehicle/vehicle group)
returning columns + rows, so the same code serves on-screen viewing, file
download and scheduled emails.
"""
import csv
import io
import logging
from datetime import datetime, timedelta, timezone
from typing import Callable, Optional
from zoneinfo import ZoneInfo

from sqlalchemy.orm import Session

from app.dataprofiler.dictionary import EVENTS, event_label
from app.models import BLETag, PortalAlert, PortalUser, UserGroup
from app.models_dataprofiler import DeviceEvent, ReportSchedule, Trip

logger = logging.getLogger(__name__)

FREQUENCIES = ("daily", "weekly", "monthly")
FORMATS = ("xlsx", "csv", "pdf")
ANALYTICS_KEY = "analytics"   # a schedule of an Analytics visual (params: page_id, visual_id)
MAX_ROWS = 50_000


# ---------------------------------------------------------------------------
# Scope helpers
# ---------------------------------------------------------------------------

def group_devices(db: Session, group_id, category: Optional[str] = None, only: Optional[set] = None) -> dict:
    """{imei: display name} for the group's devices, optionally one vehicle group (category).

    only: limit to these IMEIs (the vehicles a user may see); None means all.
    """
    tags = db.query(BLETag).filter(
        BLETag.attributes.op("->>")("_created_usergroup_id") == str(group_id)).all()
    out = {}
    for tag in tags:
        if only is not None and tag.imei not in only:
            continue
        if category and ((tag.attributes or {}).get("category") or {}).get("value") != category:
            continue
        out[tag.imei] = tag.description or tag.device_name or tag.imei
    return out


def _local(dt: Optional[datetime], tz: ZoneInfo) -> str:
    return dt.astimezone(tz).strftime("%Y-%m-%d %H:%M:%S") if dt else ""


def _dur(seconds: Optional[int]) -> str:
    s = int(seconds or 0)
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}"


# ---------------------------------------------------------------------------
# Reports: (columns, rows) for a period
# ---------------------------------------------------------------------------

def _trips_q(db, group_id, imeis, start, end):
    return (db.query(Trip).filter(Trip.usergroup_id == group_id, Trip.imei.in_(list(imeis)), Trip.real(),
                                  Trip.started_at >= start, Trip.started_at < end)
            .order_by(Trip.started_at))


def report_trips(db, group_id, devices, start, end, tz):
    from app.services import vehicle_fuel
    cols = ["Vehicle", "IMEI", "Driver", "Start", "End", "Duration", "Distance (km)", "Fuel used (L)", "Max speed (km/h)",
            "Avg speed (km/h)", "Speeding", "Harsh braking", "Harsh acceleration", "Harsh cornering",
            "Start location", "End location", "Start position", "End position"]
    rates = vehicle_fuel.group_rates(db, group_id)
    vehicle, by_id = _drivers(db, group_id)
    rows = []
    for t in _trips_q(db, group_id, devices, start, end).limit(MAX_ROWS):
        fuel = vehicle_fuel.litres(t.distance_m, rates.get(t.imei))
        rows.append([devices.get(t.imei, t.imei), t.imei, _trip_driver(t, vehicle, by_id), _local(t.started_at, tz),
                     _local(t.ended_at, tz) if t.ended_at else "in progress", _dur(t.duration_s),
                     round((t.distance_m or 0) / 1000, 2), "" if fuel is None else fuel, t.max_speed, t.avg_speed or "",
                     t.speeding_count, t.harsh_braking_count, t.harsh_acceleration_count, t.harsh_cornering_count,
                     t.start_address or "", t.end_address or "",
                     f"{t.start_latitude:.5f}, {t.start_longitude:.5f}" if t.start_latitude is not None else "",
                     f"{t.end_latitude:.5f}, {t.end_longitude:.5f}" if t.end_latitude is not None else ""])
    return cols, rows


def report_vehicle_summary(db, group_id, devices, start, end, tz):
    cols = ["Vehicle", "IMEI", "Trips", "Distance (km)", "Driving time", "Avg trip (km)",
            "Max speed (km/h)", "Speeding", "Harsh events"]
    agg: dict = {imei: [0, 0, 0, 0, 0, 0] for imei in devices}
    for t in _trips_q(db, group_id, devices, start, end):
        a = agg[t.imei]
        a[0] += 1
        a[1] += t.distance_m or 0
        a[2] += t.duration_s or 0
        a[3] = max(a[3], t.max_speed or 0)
        a[4] += t.speeding_count
        a[5] += t.harsh_braking_count + t.harsh_acceleration_count + t.harsh_cornering_count
    rows = [[name, imei, a[0], round(a[1] / 1000, 2), _dur(a[2]),
             round(a[1] / 1000 / a[0], 2) if a[0] else 0, a[3], a[4], a[5]]
            for imei, name in sorted(devices.items(), key=lambda x: x[1].lower()) for a in [agg[imei]]]
    return cols, rows


def _events(db, group_id, devices, start, end, codes=None):
    q = db.query(DeviceEvent).filter(DeviceEvent.usergroup_id == group_id, DeviceEvent.imei.in_(list(devices)),
                                     DeviceEvent.recorded_at >= start, DeviceEvent.recorded_at < end)
    if codes:
        q = q.filter(DeviceEvent.event_type.in_(codes))
    return q.order_by(DeviceEvent.recorded_at).limit(MAX_ROWS)


def report_events(db, group_id, devices, start, end, tz):
    cols = ["Vehicle", "IMEI", "Time", "Event", "Latitude", "Longitude", "Speed (km/h)", "Odometer (km)"]
    rows = []
    for e in _events(db, group_id, devices, start, end, [c for c in EVENTS if c != "parked_position"]):
        odo = (e.io or {}).get("16")
        rows.append([devices.get(e.imei, e.imei), e.imei, _local(e.recorded_at, tz),
                     event_label(e.event_type, e.event_io_id), e.latitude, e.longitude, e.speed,
                     round(odo / 1000, 1) if isinstance(odo, int) else ""])
    return cols, rows


def report_driver_behaviour(db, group_id, devices, start, end, tz):
    cols = ["Vehicle", "IMEI", "Time", "Event", "Speed (km/h)", "Detail", "Latitude", "Longitude"]
    codes = ["overspeed", "harsh_braking", "harsh_acceleration", "harsh_cornering", "idling_start", "crash"]
    rows = []
    for e in _events(db, group_id, devices, start, end, codes):
        io = e.io or {}
        detail = (f"{io.get('255')} km/h" if e.event_type == "overspeed" and io.get("255") is not None
                  else f"{io['254'] / 100:.2f} g" if isinstance(io.get("254"), int) else "")
        rows.append([devices.get(e.imei, e.imei), e.imei, _local(e.recorded_at, tz),
                     event_label(e.event_type, e.event_io_id), e.speed, detail, e.latitude, e.longitude])
    return cols, rows


def report_alerts(db, group_id, devices, start, end, tz):
    cols = ["Time", "Type", "Title", "Message", "Vehicle", "Read"]
    alerts = (db.query(PortalAlert).filter(PortalAlert.usergroup_id == str(group_id),
                                           PortalAlert.created_at >= start, PortalAlert.created_at < end)
              .order_by(PortalAlert.created_at).limit(MAX_ROWS).all())
    rows = []
    for a in alerts:
        imei = a.tag.imei if a.tag else None
        if imei and imei not in devices:
            continue
        rows.append([_local(a.created_at, tz), a.alert_type, a.title, a.message,
                     devices.get(imei, imei or ""), "yes" if a.is_read else "no"])
    return cols, rows


def _drivers(db, group_id) -> tuple:
    """({imei: the vehicle's driver}, {driver id: name}) — a trip's own driver (from a key) wins."""
    from app.services import drivers_service
    vehicle = {}
    for tag in db.query(BLETag).filter(BLETag.attributes.op("->>")("_created_usergroup_id") == str(group_id)):
        name = ((tag.attributes or {}).get("driver") or {}).get("value")
        if name:
            vehicle[tag.imei] = name
    return vehicle, drivers_service.names(db, group_id)


def _trip_driver(t, vehicle: dict, by_id: dict) -> str:
    return (by_id.get(str(t.driver_id)) if t.driver_id else None) or vehicle.get(t.imei, "")


def _speed(e) -> int:
    """The speed of an overspeed event: IO 255 (the over-speeding value), else the record's speed."""
    v = (e.io or {}).get("255")
    return int(v) if isinstance(v, (int, float)) and v > 0 else int(e.speed or 0)


def report_speeding(db, group_id, devices, start, end, tz):
    from app.services import address_service
    cols = ["Vehicle", "IMEI", "Driver", "Time", "Speed (km/h)", "Location", "Latitude", "Longitude"]
    vehicle, by_id = _drivers(db, group_id)
    events = _events(db, group_id, devices, start, end, ["overspeed"]).all()
    trips = {t.id: t for t in db.query(Trip).filter(Trip.id.in_({e.trip_id for e in events if e.trip_id}))} if events else {}
    rows = []
    for e in events:
        t = trips.get(e.trip_id)
        driver = _trip_driver(t, vehicle, by_id) if t else vehicle.get(e.imei, "")
        rows.append([devices.get(e.imei, e.imei), e.imei, driver, _local(e.recorded_at, tz), _speed(e),
                     address_service.cached(db, e.latitude, e.longitude) or "",
                     e.latitude, e.longitude])
    return cols, rows


def report_speeding_summary(db, group_id, devices, start, end, tz):
    cols = ["Vehicle", "IMEI", "Assigned driver", "Speeding events", "Top speed (km/h)", "Average speeding speed (km/h)",
            "Distance (km)", "Speeding per 100 km"]
    vehicle, _ = _drivers(db, group_id)
    agg = {imei: [0, 0, 0] for imei in devices}           # count, top, sum of speeds
    for e in _events(db, group_id, devices, start, end, ["overspeed"]):
        a, v = agg[e.imei], _speed(e)
        a[0] += 1
        a[1] = max(a[1], v)
        a[2] += v
    km = {imei: 0.0 for imei in devices}
    for t in _trips_q(db, group_id, devices, start, end):
        km[t.imei] += (t.distance_m or 0) / 1000
    rows = []
    for imei, name in sorted(devices.items(), key=lambda x: (-agg[x[0]][0], x[1].lower())):
        n, top, total = agg[imei]
        rows.append([name, imei, vehicle.get(imei, ""), n, top or "", round(total / n) if n else "", round(km[imei], 1),
                     round(n / km[imei] * 100, 2) if km[imei] else ""])
    return cols, rows


def _after_hours(db, group_id, devices, start, end, tz):
    """[(trip, inside s, outside s, profile summary)] using each vehicle's time profile (or the default)."""
    from app.services import time_profiles
    by_id, default = time_profiles.profiles(db, group_id)
    attrs = {tag.imei: tag.attributes or {} for tag in
             db.query(BLETag).filter(BLETag.attributes.op("->>")("_created_usergroup_id") == str(group_id))}
    out = []
    for t in _trips_q(db, group_id, devices, start, end):
        windows = time_profiles.for_vehicle(attrs.get(t.imei), by_id, default)
        ends = t.ended_at or t.last_event_at
        inside, outside = time_profiles.split(time_profiles.utc(t.started_at), time_profiles.utc(ends), windows, tz)
        out.append((t, inside, outside, time_profiles.describe(windows)))
    return out


def report_after_hours(db, group_id, devices, start, end, tz):
    cols = ["Vehicle", "IMEI", "Working hours", "Trips", "After-hours trips", "Distance (km)", "After-hours distance (km)",
            "Driving time", "After-hours driving time", "After hours (% of distance)"]
    agg = {imei: [0, 0, 0.0, 0.0, 0, 0, ""] for imei in devices}
    for t, inside, outside, profile in _after_hours(db, group_id, devices, start, end, tz):
        a, km = agg[t.imei], (t.distance_m or 0) / 1000
        total = inside + outside
        a[0] += 1
        a[1] += 1 if outside > 0 else 0
        a[2] += km
        a[3] += km * outside / total if total else 0
        a[4] += t.duration_s or 0
        a[5] += outside
        a[6] = profile
    rows = []
    for imei, name in sorted(devices.items(), key=lambda x: (-agg[x[0]][3], x[1].lower())):
        n, n_out, km, km_out, secs, secs_out, profile = agg[imei]
        rows.append([name, imei, profile, n, n_out, round(km, 1), round(km_out, 1), _dur(secs), _dur(secs_out),
                     round(km_out / km * 100) if km else ""])
    return cols, rows


def report_after_hours_trips(db, group_id, devices, start, end, tz):
    cols = ["Vehicle", "IMEI", "Driver", "Start", "End", "Duration", "After-hours time", "Distance (km)",
            "After-hours distance (km)", "Start location", "End location"]
    vehicle, by_id = _drivers(db, group_id)
    rows = []
    for t, inside, outside, _ in _after_hours(db, group_id, devices, start, end, tz):
        if outside <= 0:
            continue
        km, total = (t.distance_m or 0) / 1000, inside + outside
        rows.append([devices.get(t.imei, t.imei), t.imei, _trip_driver(t, vehicle, by_id), _local(t.started_at, tz),
                     _local(t.ended_at, tz) if t.ended_at else "in progress", _dur(t.duration_s), _dur(outside),
                     round(km, 2), round(km * outside / total, 2) if total else 0, t.start_address or "", t.end_address or ""])
    return cols, rows


def report_licences(db, group_id, devices, start, end, tz):
    """Every vehicle's licence: start, expiry and days left (the period isn't used)."""
    from app.services import licence_service
    cols = ["Vehicle", "IMEI", "Registration", "Licence start", "Licence expires", "Days left", "Status"]
    today = datetime.now(tz).date()
    rows = []
    for tag in db.query(BLETag).filter(BLETag.imei.in_(list(devices))):
        attrs = tag.attributes or {}
        st = licence_service.status(licence_service.licence_start(attrs), today)
        label = {"valid": "Valid", "expiring": "Expires within 7 days", "expired": "Expired", "unknown": "No licence date"}[st["state"]]
        rows.append([devices[tag.imei], tag.imei, ((attrs.get("registration") or {}).get("value") or ""), st["start"] or "",
                     st["expires"] or "", st["days_left"] if st["days_left"] is not None else "", label])
    rows.sort(key=lambda r: (r[5] == "", r[5] if r[5] != "" else 0))
    return cols, rows


# key → (title, category, description, generator)
REPORTS: dict[str, tuple[str, str, str, Callable]] = {
    "trips":            ("Trip detail", "Trips", "Every trip with duration, distance, speeds and driving events.", report_trips),
    "vehicle_summary":  ("Vehicle summary", "Utilisation", "Trips, distance and driving time per vehicle.", report_vehicle_summary),
    "driver_behaviour": ("Driving behaviour", "Driving", "Speeding, harsh driving, idling and crash events.", report_driver_behaviour),
    "speeding":         ("Speeding detail", "Driving", "Every speeding event: vehicle, driver, time, speed and where.", report_speeding),
    "speeding_summary": ("Speeding summary", "Driving", "Speeding events, top speed and speeding per 100 km per vehicle.", report_speeding_summary),
    "after_hours":      ("After-hours usage", "Utilisation", "Trips, distance and driving time outside working hours per vehicle (Configuration → Time Profiles).", report_after_hours),
    "after_hours_trips": ("After-hours trips", "Utilisation", "Every trip that was partly or wholly outside working hours.", report_after_hours_trips),
    "licences":         ("Vehicle licences", "Vehicles", "Each vehicle's licence start, expiry and days left.", report_licences),
    "events":           ("Event detail", "Events", "All decoded device events (excluding parked updates).", report_events),
    "alerts":           ("Alert history", "Alerts", "Notifications raised for your vehicles.", report_alerts),
}


def catalogue() -> list[dict]:
    return [{"key": k, "title": t, "category": c, "description": d} for k, (t, c, d, _) in REPORTS.items()]


def generate(db: Session, group_id, key: str, start: datetime, end: datetime,
             imei: Optional[str] = None, category: Optional[str] = None,
             tz_name: str = "Africa/Johannesburg", only: Optional[set] = None) -> tuple[list, list]:
    if key not in REPORTS:
        raise KeyError(key)
    devices = group_devices(db, group_id, category, only)
    if imei:
        devices = {k: v for k, v in devices.items() if k == imei}
    return REPORTS[key][3](db, group_id, devices, start, end, ZoneInfo(tz_name))


def _logo_image(data_url: Optional[str]):
    """An openpyxl image from a PNG/JPEG data URL (SVG isn't supported in Excel files), or None."""
    if not data_url or not data_url.startswith(("data:image/png", "data:image/jpeg", "data:image/jpg")):
        return None
    try:
        import base64
        from openpyxl.drawing.image import Image
        img = Image(io.BytesIO(base64.b64decode(data_url.split(",", 1)[1])))
        scale = min(1.0, 60 / img.height, 220 / img.width)   # fit 220 × 60 px
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        return img
    except Exception:
        logger.warning("Report logo could not be used", exc_info=True)
        return None


_PDF_CHARS = str.maketrans({"–": "-", "—": "-", "·": "-", "’": "'", "‘": "'", "“": '"', "”": '"', "…": "...",
                            "→": "->", "≥": ">=", "≤": "<=", "×": "x", "÷": "/", "€": "EUR"})


def _latin(value) -> str:
    """Text the PDF's built-in font can show."""
    return str("" if value is None else value).translate(_PDF_CHARS).encode("latin-1", "replace").decode("latin-1")


def _text(v) -> str:
    """A cell value as text: thousands separated, floats without trailing zeros."""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, float):
        return f"{v:,.2f}".rstrip("0").rstrip(".")
    if isinstance(v, int):
        return f"{v:,}"
    return _latin(v)


def _rgb(hex_colour: Optional[str]):
    h = (hex_colour or "").lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4)) if len(h) == 6 else None


def _pdf(title: str, cols: list, rows: list, subtitle: Optional[str], logo: Optional[str],
         styles: Optional[list] = None, orientation: str = "landscape") -> bytes:
    """An A4 table: logo, title and subtitle, then the rows with the header repeated on every page."""
    from fpdf import FPDF
    pdf = FPDF(orientation="P" if orientation == "portrait" else "L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    if logo and logo.startswith(("data:image/png", "data:image/jpeg", "data:image/jpg")):
        try:
            import base64
            pdf.image(io.BytesIO(base64.b64decode(logo.split(",", 1)[1])), x=10, y=8, h=14)
            pdf.set_y(24)
        except Exception:
            logger.warning("Report logo could not be used in the PDF", exc_info=True)
    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 8, _latin(title), new_x="LMARGIN", new_y="NEXT")
    if subtitle:
        pdf.set_font("Helvetica", "", 9)
        pdf.set_text_color(110, 110, 110)
        pdf.cell(0, 6, _latin(subtitle), new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0, 0, 0)
    pdf.ln(2)
    if not cols:
        return bytes(pdf.output())
    # Column widths from the measured text (values in full; headers wrap); the table spans the page
    from fpdf.fonts import FontFace
    sample = rows[:300]
    pdf.set_font("Helvetica", "", 7.5)
    body = [max([0.0] + [pdf.get_string_width(_text(r[i])) for r in sample if i < len(r)]) for i in range(len(cols))]
    pdf.set_font("Helvetica", "B", 7.5)
    head = [pdf.get_string_width(_latin(c)) for c in cols]
    want = [max(b, h * 0.55, 10) + 3 for b, h in zip(body, head)]
    numeric = [all(isinstance(r[i], (int, float)) or r[i] in ("", None) for r in sample if i < len(r)) for i in range(len(cols))]
    pdf.set_font("Helvetica", "", 7.5)
    with pdf.table(col_widths=tuple(want), text_align=tuple("RIGHT" if n else "LEFT" for n in numeric),
                   line_height=4.2, borders_layout="NONE", cell_fill_color=(246, 248, 251), cell_fill_mode="ROWS",
                   headings_style=FontFace(emphasis="BOLD", fill_color=(230, 236, 245)), repeat_headings=1,
                   padding=(1.2, 1.5)) as table:
        heading = table.row()
        for c in cols:
            heading.cell(_latin(c))
        for n, r in enumerate(rows):
            row = table.row()
            style = styles[n] if styles and n < len(styles) else None
            for i in range(len(cols)):
                bg, fg = ((style["cells"].get(i) or style["row"] or (None, None)) if style else (None, None))
                face = FontFace(fill_color=_rgb(bg), color=_rgb(fg)) if (bg or fg) else None
                row.cell(_text(r[i] if i < len(r) else ""), style=face)
    pdf.set_font("Helvetica", "", 7)
    pdf.set_text_color(130, 130, 130)
    pdf.cell(0, 6, f"{len(rows)} rows", new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def analytics_table(db: Session, group_id, page_id: str, visual_id: str, start: datetime, end: datetime, tz_name: str,
                    only: Optional[set] = None, category: Optional[str] = None, imei: Optional[str] = None) -> tuple:
    """(title, columns, rows) of an Analytics visual for a period — for scheduled e-mails and downloads."""
    from uuid import UUID
    from app.models import AnalyticsExpression, AnalyticsPage
    from app.services import analytics_service as an
    page = db.query(AnalyticsPage).filter(AnalyticsPage.id == UUID(str(page_id)), AnalyticsPage.usergroup_id == group_id).first()
    visual = next((v for v in (page.visuals or []) if v.get("id") == visual_id), None) if page else None
    if not visual:
        raise KeyError("The Analytics chart no longer exists")
    expressions = {str(e.id): (e.name, e.expression, e.decimals, e.unit)
                   for e in db.query(AnalyticsExpression).filter(AnalyticsExpression.usergroup_id == group_id)}
    devices = an.devices_for(db, group_id, only, visual.get("category") or category, visual.get("imei") or imei)
    result = an.query(db, group_id, devices, start, end, ZoneInfo(tz_name), dimension=visual.get("dimension", ""),
                      measures=visual.get("measures") or [], expressions=expressions, sort=visual.get("sort"),
                      descending=visual.get("descending", True), limit=visual.get("limit"), place_id=visual.get("place_id"))
    return f"{page.name} - {visual.get('title') or 'Chart'}", *result_table(result)


def result_table(result: dict) -> tuple:
    """(columns, rows) of an analytics result: the row label, its text columns, then each value (rounded)."""
    info = result.get("info_columns") or []
    record = result.get("kind") == "record"
    cols = ([] if record or result["dimension"] == "" else [result["dimension_label"]]) + [c["label"] for c in info] + [
        f"{c['label']} ({c['unit']})" if c["unit"] else c["label"] for c in result["columns"]]
    rows = []
    for r in result["rows"]:
        values = []
        for c in result["columns"]:
            v = r["values"].get(c["key"])
            values.append("" if v is None else (round(v) if c["decimals"] == 0 else round(v, c["decimals"])))
        rows.append(([] if record or result["dimension"] == "" else [r["label"]])
                    + [(r.get("info") or {}).get(c["key"], "") for c in info] + values)
    return cols, rows


def to_file(title: str, cols: list, rows: list, fmt: str,
            subtitle: Optional[str] = None, logo: Optional[str] = None, styles: Optional[list] = None,
            orientation: str = "landscape") -> tuple[bytes, str, str]:
    """(content, mime type, extension). Excel files get a heading (logo, title, subtitle) when given.

    styles: per row {"row": (bg, fg) | None, "cells": {col index: (bg, fg)}} — conditional colours (Excel, PDF).
    """
    if fmt == "pdf":
        return _pdf(title, cols, rows, subtitle, logo, styles, orientation), "application/pdf", "pdf"
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(cols)
        writer.writerows(rows)
        return buf.getvalue().encode("utf-8-sig"), "text/csv", "csv"
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    header_row = 1
    if subtitle or logo:
        image = _logo_image(logo)
        if image is not None:
            ws.add_image(image, "A1")
            ws.row_dimensions[1].height = image.height * 0.75 + 4   # px → points
            ws.append([])
        ws.append([title])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=14)
        if subtitle:
            ws.append([subtitle])
            ws.cell(row=ws.max_row, column=1).font = Font(color="666666")
        ws.append([])
        header_row = ws.max_row + 1
    ws.append(cols)
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    from openpyxl.styles import PatternFill
    for n, row in enumerate(rows):
        ws.append(row)
        style = (styles or [None] * len(rows))[n] if styles else None
        if style:
            for i in range(len(cols)):
                bg, fg = style["cells"].get(i) or style["row"] or (None, None)
                cell = ws.cell(row=ws.max_row, column=i + 1)
                if bg:
                    cell.fill = PatternFill("solid", fgColor=bg.lstrip("#").upper())
                if fg:
                    cell.font = Font(color=fg.lstrip("#").upper())
    for i, col in enumerate(cols, start=1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in rows[:500]]) + 2
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = min(width, 50)
    ws.freeze_panes = f"A{header_row + 1}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"


# ---------------------------------------------------------------------------
# Scheduling
# ---------------------------------------------------------------------------

def period_for(frequency: str, run_at: datetime, tz: ZoneInfo) -> tuple[datetime, datetime]:
    """The whole previous day / week (Mon–Sun) / month before run_at, in tz, as UTC."""
    local = run_at.astimezone(tz).replace(hour=0, minute=0, second=0, microsecond=0)
    if frequency == "daily":
        start, end = local - timedelta(days=1), local
    elif frequency == "weekly":
        end = local - timedelta(days=local.weekday())
        start = end - timedelta(days=7)
    else:
        end = local.replace(day=1)
        start = (end - timedelta(days=1)).replace(day=1)
    return start.astimezone(timezone.utc), end.astimezone(timezone.utc)


def next_run(frequency: str, send_hour: int, tz_name: str, after: datetime) -> datetime:
    """First send time strictly after `after`: daily at send_hour, weekly Mondays, monthly on the 1st."""
    tz = ZoneInfo(tz_name)
    local = after.astimezone(tz)
    candidate = local.replace(hour=send_hour, minute=0, second=0, microsecond=0)
    while True:
        if candidate > local and (frequency == "daily"
                                  or (frequency == "weekly" and candidate.weekday() == 0)
                                  or (frequency == "monthly" and candidate.day == 1)):
            return candidate.astimezone(timezone.utc)
        candidate += timedelta(days=1)


def _creator_scope(db: Session, schedule: ReportSchedule) -> Optional[set]:
    """The vehicles the schedule's creator may see (None = all), so emailed reports respect security groups."""
    from app.services import access_service
    creator = db.get(PortalUser, schedule.created_by) if schedule.created_by else None
    if creator is None:
        return None
    return access_service.for_user(creator, db).visible_imeis(db)


def _run_custom(db: Session, schedule: ReportSchedule, email_service, start: datetime, end: datetime, tz: ZoneInfo) -> str:
    """A scheduled custom report: generate it (kept in the report's history) and e-mail the file."""
    from uuid import UUID
    from app.models_reports import CustomReport
    from app.routes_custom_reports import generate
    params = schedule.params or {}
    report = db.query(CustomReport).filter(CustomReport.id == UUID(params["report_id"]),
                                           CustomReport.usergroup_id == schedule.usergroup_id).first()
    if not report:
        return "failed: the report was deleted"
    values = {**(params.get("values") or {}), "date_range": (start, end)}
    run = generate(db, report, values, start, end, only=_creator_scope(db, schedule), created_by="schedule",
                   schedule_id=schedule.id, file_name=params.get("file_name"), compress=bool(params.get("compress")))
    if run.status != "done":
        return f"failed: {run.error}"
    period = f"{start.astimezone(tz):%d %b %Y}" + (
        "" if schedule.frequency == "daily" else f" – {(end - timedelta(seconds=1)).astimezone(tz):%d %b %Y}")
    name = schedule.name or report.title
    subject = params.get("subject") or f"{name} — {period}"
    text_body = params.get("body") or f"Your {schedule.frequency} report \"{name}\" for {period} is attached ({run.rows} rows)."
    html_body = "".join(f"<p>{line}</p>" for line in text_body.splitlines() if line.strip())
    sent = sum(email_service.send_report_email(to, subject, text_body, html_body, run.file_name, run.content, run.mime)
               for to in schedule.recipients or [])
    return f"sent to {sent}/{len(schedule.recipients or [])} recipients ({run.rows} rows)"


def run_schedule(db: Session, schedule: ReportSchedule, email_service, now: datetime) -> str:
    tz = ZoneInfo(schedule.timezone)
    start, end = period_for(schedule.frequency, now, tz)
    params = schedule.params or {}
    if schedule.report_key == "custom":
        return _run_custom(db, schedule, email_service, start, end, tz)
    if schedule.report_key == ANALYTICS_KEY:
        title, cols, rows = analytics_table(db, schedule.usergroup_id, params.get("page_id"), params.get("visual_id"), start, end,
                                            schedule.timezone, only=_creator_scope(db, schedule),
                                            category=params.get("category"), imei=params.get("imei"))
    else:
        title = REPORTS[schedule.report_key][0]
        cols, rows = generate(db, schedule.usergroup_id, schedule.report_key, start, end,
                              imei=params.get("imei"), category=params.get("category"), tz_name=schedule.timezone,
                              only=_creator_scope(db, schedule))
    period = f"{start.astimezone(tz):%d %b %Y}" + (
        "" if schedule.frequency == "daily" else f" – {(end - timedelta(seconds=1)).astimezone(tz):%d %b %Y}")
    group = db.get(UserGroup, schedule.usergroup_id)
    content, mime, ext = to_file(title, cols, rows, schedule.file_format,
                                 subtitle=f"{group.name if group else ''} · {period}".strip(" ·"),
                                 logo=group.report_logo_data if group else None)
    name = schedule.name or title
    filename = f"{name.replace(' ', '_')}_{start.astimezone(tz):%Y-%m-%d}.{ext}"
    subject = f"{name} — {period}"
    text_body = f"Your {schedule.frequency} report \"{name}\" for {period} is attached ({len(rows)} rows)."
    html_body = (f"<p>Your {schedule.frequency} report <strong>{name}</strong> for {period} is attached "
                 f"({len(rows)} rows).</p><p style='color:#888;font-size:12px'>Manage scheduled reports under "
                 f"Reports in {email_service.brand_name}.</p>")
    sent = sum(email_service.send_report_email(to, subject, text_body, html_body, filename, content, mime)
               for to in schedule.recipients or [])
    return f"sent to {sent}/{len(schedule.recipients or [])} recipients ({len(rows)} rows)"


def run_due(db: Session, email_service=None, now: Optional[datetime] = None) -> int:
    """Send every schedule that is due, then move it to its next run. Commits per schedule."""
    if email_service is None:
        from app.services.email_service import EmailService
        email_service = EmailService()
    now = now or datetime.now(timezone.utc)
    due = db.query(ReportSchedule).filter(ReportSchedule.is_active == True, ReportSchedule.next_run_at <= now).all()
    for schedule in due:
        try:
            schedule.last_status = run_schedule(db, schedule, email_service, now)
        except Exception as exc:
            logger.exception("Scheduled report %s failed", schedule.id)
            schedule.last_status = f"failed: {exc.__class__.__name__}"
        schedule.last_run_at = now
        schedule.next_run_at = next_run(schedule.frequency, schedule.send_hour, schedule.timezone, now)
        db.commit()
    return len(due)
